import os
import re
import struct
import sys
import unittest
import zlib
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.append(str(Path(__file__).resolve().parents[1]))

if "fitz" not in sys.modules:
    fitz_module = type(sys)("fitz")

    class _Tools:
        @staticmethod
        def mupdf_display_errors(_value):
            return None

    def _not_available(*args, **kwargs):
        raise NotImplementedError

    fitz_module.TOOLS = _Tools()
    fitz_module.open = _not_available
    sys.modules["fitz"] = fitz_module

if "psycopg2" not in sys.modules:
    psycopg2_module = type(sys)("psycopg2")
    psycopg2_pool_module = type(sys)("psycopg2.pool")
    psycopg2_extras_module = type(sys)("psycopg2.extras")

    def _connect(*args, **kwargs):
        raise NotImplementedError

    class _SimpleConnectionPool:
        def __init__(self, *args, **kwargs):
            pass

    class _ThreadedConnectionPool(_SimpleConnectionPool):
        pass

    psycopg2_module.connect = _connect
    psycopg2_module.Error = Exception
    psycopg2_pool_module.SimpleConnectionPool = _SimpleConnectionPool
    psycopg2_pool_module.ThreadedConnectionPool = _ThreadedConnectionPool
    psycopg2_extras_module.RealDictCursor = object
    psycopg2_module.pool = psycopg2_pool_module
    psycopg2_module.extras = psycopg2_extras_module
    sys.modules["psycopg2"] = psycopg2_module
    sys.modules["psycopg2.pool"] = psycopg2_pool_module
    sys.modules["psycopg2.extras"] = psycopg2_extras_module

os.environ.setdefault("DASHSCOPE_API_KEY", "test-key")
os.environ.setdefault("OSS_BUCKET", "test-bucket")
os.environ.setdefault("PG_HOST", "localhost")
os.environ.setdefault("PG_USER", "postgres")
os.environ.setdefault("PG_PASSWORD", "postgres")
os.environ.setdefault("MILVUS_HOST", "localhost")
os.environ.setdefault("OSS_ACCESS_KEY_ID", "key")
os.environ.setdefault("OSS_ACCESS_KEY_SECRET", "secret")

from app.core.exceptions import ValidationError
from app.services import doc_image_parser, document_service, job_service
from app.services.chunk_splitter import split_parent_child_text, split_text_with_metadata


def _chunk(tag: bytes, payload: bytes) -> bytes:
    checksum = zlib.crc32(tag + payload) & 0xFFFFFFFF
    return struct.pack("!I", len(payload)) + tag + payload + struct.pack("!I", checksum)


def _build_png(width: int = 48, height: int = 48) -> bytes:
    rows = []
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            row.extend(((x * 5) % 256, (y * 3) % 256, ((x + y) * 7) % 256))
        rows.append(bytes(row))
    payload = zlib.compress(b"".join(rows), level=0)
    header = struct.pack("!2I5B", width, height, 8, 2, 0, 0, 0)
    return b"".join(
        [
            b"\x89PNG\r\n\x1a\n",
            _chunk(b"IHDR", header),
            _chunk(b"IDAT", payload),
            _chunk(b"IEND", b""),
        ]
    )


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)

    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    run.append(run_props)

    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _build_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("全品牌通用知识", level=1)

    paragraph = doc.add_paragraph("查看资料：")
    _add_hyperlink(paragraph, "飞书资料", "https://example.com/doc")
    paragraph.add_run("，请按指引执行。")

    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "审批类别"
    table.rows[0].cells[1].text = "适用场景"
    table.rows[1].cells[0].text = "开店申请"
    table.rows[1].cells[1].text = "新加盟商首店"
    table.rows[2].cells[0].text = "人员入企"
    table.rows[2].cells[1].text = "门店培训"

    doc.add_paragraph("示意图如下：")
    run = doc.add_paragraph().add_run()
    run.add_picture(BytesIO(_build_png()), width=None, height=None)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FAQ"
    sheet.append(["问题", "答案"])
    sheet.append(["店长端无法打开", "请检查网络和版本"])
    sheet.append(["开店审核慢怎么办", "联系运营同学协助排查"])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class DocxPipelineTests(unittest.TestCase):
    def test_parent_child_text_chunks_keep_parent_context(self):
        text = "第一段说明开店流程。" * 80 + "\n\n" + "第二段说明人员入企。" * 80

        chunks = split_parent_child_text(
            text,
            parent_chunk_size=300,
            child_chunk_size=120,
            chunk_overlap=20,
            base_metadata={"file_name": "guide.txt", "source": "txt"},
            parent_id_prefix="job-a",
        )

        self.assertGreater(len(chunks), 2)
        self.assertEqual(chunks[0]["metadata"]["chunk_strategy"], "parent_child")
        self.assertEqual(chunks[0]["metadata"]["parent_id"], "job-a-parent-0")
        self.assertIn("parent_content", chunks[0]["metadata"])
        self.assertLessEqual(len(chunks[0]["content"]), 180)

    def test_flat_text_strategy_remains_available(self):
        chunks = split_text_with_metadata(
            "A" * 260,
            chunk_size=100,
            chunk_overlap=0,
            base_metadata={"source": "txt"},
            chunk_strategy="flat",
        )

        self.assertEqual(chunks[0]["metadata"]["chunk_strategy"], "flat")
        self.assertNotIn("parent_content", chunks[0]["metadata"])

    def test_parse_text_mode_docx_preserves_markdown_tables_and_hyperlinks(self):
        chunks, _ = job_service._parse_text_mode(
            file_content=_build_docx_bytes(),
            file_name="guide.docx",
            job_id="job-text-1",
            chunk_size=2_000,
            chunk_overlap=0,
        )

        content = "\n".join(chunk["content"] for chunk in chunks)

        self.assertIn("全品牌通用知识", content)
        self.assertIn("飞书资料 (https://example.com/doc)", content)
        self.assertIn("| 审批类别 | 适用场景 |", content)
        self.assertIn("| --- | --- |", content)
        self.assertIn("| 开店申请 | 新加盟商首店 |", content)
        self.assertNotIn("<<IMAGE:", content)

    def test_parse_text_mode_accepts_parent_child_config(self):
        chunks, images = job_service._parse_text_mode(
            file_content=("开店流程。" * 120).encode("utf-8"),
            file_name="guide.txt",
            job_id="job-parent-child",
            chunk_size=500,
            chunk_overlap=20,
            parent_chunk_size=300,
            child_chunk_size=120,
            chunk_strategy="parent_child",
        )

        self.assertEqual(images, [])
        self.assertGreater(len(chunks), 1)
        self.assertEqual(chunks[0]["metadata"]["chunk_strategy"], "parent_child")
        self.assertIn("parent_content", chunks[0]["metadata"])

    def test_parse_word_docx_preserves_text_structure_and_inserts_placeholders(self):
        original_upload = doc_image_parser._upload_image
        doc_image_parser._upload_image = (
            lambda image_bytes, ext, collection, file_name, chunk_id: f"oss://{chunk_id}.{ext}"
        )
        try:
            chunks, image_records = doc_image_parser.parse_word(
                file_content=_build_docx_bytes(),
                job_id="job-1",
                collection="kb-demo",
                file_name="guide.docx",
                chunk_size=500,
                chunk_overlap=20,
                parent_chunk_size=900,
                child_chunk_size=260,
                chunk_strategy="parent_child",
            )
        finally:
            doc_image_parser._upload_image = original_upload

        content = "\n".join(chunk["content"] for chunk in chunks)

        self.assertIn("飞书资料 (https://example.com/doc)", content)
        self.assertIn("| 审批类别 | 适用场景 |", content)
        self.assertEqual(chunks[0]["metadata"]["chunk_strategy"], "parent_child")
        self.assertIn("parent_content", chunks[0]["metadata"])
        self.assertEqual(len(image_records), 1)
        self.assertRegex(content, r"<<IMAGE:[0-9a-f]{8}>>")
        chunk_ids = {chunk["chunk_id"] for chunk in chunks}
        self.assertIn(image_records[0]["chunk_id"], chunk_ids)
        self.assertTrue(any(
            image_records[0]["placeholder"] in chunk["content"]
            or image_records[0]["placeholder"] in chunk["metadata"].get("parent_content", "")
            for chunk in chunks
        ))

    def test_parse_text_mode_adds_retrieval_bucket_for_manual_docs(self):
        chunks, _ = job_service._parse_text_mode(
            file_content=("开店流程说明。" * 120).encode("utf-8"),
            file_name="guide.txt",
            job_id="job-manual-bucket",
            chunk_size=500,
            chunk_overlap=20,
            parent_chunk_size=300,
            child_chunk_size=120,
            chunk_strategy="parent_child",
            chunk_profile="smart_mix",
        )

        self.assertGreater(len(chunks), 1)
        self.assertEqual(chunks[0]["metadata"]["chunk_strategy"], "parent_child")
        self.assertEqual(chunks[0]["metadata"]["retrieval_bucket"], "manual")
        self.assertEqual(chunks[0]["metadata"]["chunk_profile"], "smart_mix")
        self.assertIn("parent_content", chunks[0]["metadata"])

    def test_parse_text_mode_routes_faq_named_docs_to_flat_chunks(self):
        chunks, _ = job_service._parse_text_mode(
            file_content=("店长端无法打开，请检查网络后重试。" * 80).encode("utf-8"),
            file_name="merchant_qa.txt",
            job_id="job-faq-flat",
            chunk_size=500,
            chunk_overlap=20,
            parent_chunk_size=300,
            child_chunk_size=120,
            chunk_strategy="parent_child",
            chunk_profile="smart_mix",
        )

        self.assertGreater(len(chunks), 1)
        self.assertEqual(chunks[0]["metadata"]["chunk_strategy"], "flat")
        self.assertEqual(chunks[0]["metadata"]["retrieval_bucket"], "faq")
        self.assertEqual(chunks[0]["metadata"]["chunk_profile"], "smart_mix")
        self.assertNotIn("parent_content", chunks[0]["metadata"])

    def test_parse_text_mode_excel_chunks_are_tagged_as_faq(self):
        chunks, excel_images = job_service._parse_text_mode(
            file_content=_build_xlsx_bytes(),
            file_name="tickets_qa.xlsx",
            job_id="job-excel-faq",
            chunk_size=500,
            chunk_overlap=20,
            chunk_profile="smart_mix",
            excel_rows_per_chunk=1,
        )

        self.assertEqual(excel_images, [])
        self.assertGreaterEqual(len(chunks), 2)
        self.assertEqual(chunks[0]["metadata"]["retrieval_bucket"], "faq")
        self.assertEqual(chunks[0]["metadata"]["chunk_profile"], "smart_mix")
        self.assertIn(chunks[0]["metadata"]["chunk_strategy"], {"excel_rows", "excel_image_rows"})

    def test_parse_word_adds_retrieval_bucket_for_manual_docs(self):
        original_upload = doc_image_parser._upload_image
        doc_image_parser._upload_image = (
            lambda image_bytes, ext, collection, file_name, chunk_id: f"oss://{chunk_id}.{ext}"
        )
        try:
            chunks, _ = doc_image_parser.parse_word(
                file_content=_build_docx_bytes(),
                job_id="job-word-bucket",
                collection="kb-demo",
                file_name="guide.docx",
                chunk_size=500,
                chunk_overlap=20,
                parent_chunk_size=900,
                child_chunk_size=260,
                chunk_strategy="parent_child",
                chunk_profile="smart_mix",
            )
        finally:
            doc_image_parser._upload_image = original_upload

        self.assertGreaterEqual(len(chunks), 1)
        self.assertEqual(chunks[0]["metadata"]["retrieval_bucket"], "manual")
        self.assertEqual(chunks[0]["metadata"]["chunk_profile"], "smart_mix")

    def test_image_mode_pipeline_does_not_require_excel_image_data(self):
        class FakeJobRepo:
            def __init__(self):
                self.statuses = []

            def update_status(self, job_id, status, **kwargs):
                self.statuses.append((status, kwargs))

        class FakeFileRepo:
            def __init__(self):
                self.statuses = []

            def update_status(self, file_id, status, error_msg=None):
                self.statuses.append((status, error_msg))

        class FakeChunkRepo:
            def __init__(self):
                self.inserted = False

            def bulk_insert_with_ids(self, job_id, file_name, chunks):
                self.inserted = True

        class FakeChunkImageRepo:
            def bulk_insert(self, image_records):
                pass

        class FakeKbRepo:
            def get_by_id(self, kb_id):
                return {"metadata_fields": []}

        job_repo = FakeJobRepo()
        file_repo = FakeFileRepo()
        chunk_repo = FakeChunkRepo()

        def fake_parse_image_mode(**kwargs):
            return ([{"chunk_id": "chunk-1", "content": "hello", "metadata": {}}], [])

        with (
            patch.object(job_service, "get_job_repository", return_value=job_repo),
            patch.object(job_service, "get_file_repository", return_value=file_repo),
            patch.object(job_service, "get_chunk_repository", return_value=chunk_repo),
            patch.object(job_service, "get_chunk_image_repository", return_value=FakeChunkImageRepo()),
            patch.object(job_service, "_download_file", return_value=b"docx-bytes"),
            patch.object(job_service, "_parse_image_mode", side_effect=fake_parse_image_mode),
            patch("app.services.chunk_service._delete_job_images_from_oss", return_value=None),
            patch("app.db.get_kb_repository", return_value=FakeKbRepo()),
        ):
            import asyncio

            asyncio.run(job_service.run_job_pipeline(
                job_id="job-image-1",
                file_id="file-image-1",
                kb_id="kb-image-1",
                kb_name="kb-demo",
                file_name="guide.docx",
                oss_key="kb-demo/guide.docx",
                image_mode=True,
                chunk_size=2000,
                chunk_overlap=0,
                image_dpi=144,
            ))

        self.assertTrue(chunk_repo.inserted)
        self.assertNotIn("error", [status for status, _ in job_repo.statuses])
        self.assertEqual(job_repo.statuses[-1][0], "chunked")

    def test_validate_file_rejects_legacy_doc_with_conversion_message(self):
        with self.assertRaisesRegex(ValidationError, r"\.docx"):
            document_service.validate_file("legacy.doc", 128)


if __name__ == "__main__":
    unittest.main()
